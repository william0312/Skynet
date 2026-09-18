# -*- coding: latin-1 -*-
import math
import pandas as pd
from datetime import datetime
import pytz
import smtplib
from email.message import EmailMessage

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.styles import PatternFill
from openpyxl.styles import Font

def procesar_metricas_matricula(df_matricula, vel_subida, vel_bajada):
  """Procesa de forma iterativa los rangos del 6 al 20 guardando las métricas en un diccionario anidado sin usar globals()."""
  metricas = {}

  for r in range(6, 21):
    df_rango = df_matricula[df_matricula['Rango'] == r]
    num_filas = len(df_rango)

    pos_p5 = math.ceil(num_filas * 0.05)
    pos_p95 = math.ceil(num_filas * 0.95)

    # Inicialización por defecto en el diccionario interno
    metricas[r] = {
        'pos_p5_sub': pos_p5,
        'val_p5_sub': 0,
        'res_p5_sub': 'NO APLICA',
        'pos_p95_sub': pos_p95,
        'val_p95_sub': 0,
        'res_p95_sub': 'NO APLICA',
        'pos_p5_baj': pos_p5,
        'val_p5_baj': 0,
        'res_p5_baj': 'NO APLICA',
        'pos_p95_baj': pos_p95,
        'val_p95_baj': 0,
        'res_p95_baj': 'NO APLICA',
    }

    # Evaluación y cálculo si existen registros en el rango
    if num_filas > 0:
      # --- Lógica de Subida ---
      df_subida = df_rango.sort_values(by='Velocidad de subida [Kbps]')

      fila_subida_p5 = df_subida.iloc[pos_p5 - 1]
      val_p5_sub = fila_subida_p5['Velocidad de subida [Kbps]']
      metricas[r]['val_p5_sub'] = val_p5_sub
      metricas[r]['res_p5_sub'] = (
          'CUMPLE' if val_p5_sub >= vel_subida else 'NO CUMPLE'
      )

      fila_subida_p95 = df_subida.iloc[pos_p95 - 1]
      val_p95_sub = fila_subida_p95['Velocidad de subida [Kbps]']
      metricas[r]['val_p95_sub'] = val_p95_sub
      metricas[r]['res_p95_sub'] = (
          'CUMPLE' if val_p95_sub >= vel_subida else 'NO CUMPLE'
      )

      # --- Lógica de Bajada ---
      df_bajada = df_rango.sort_values(by='Velocidad de bajada [Kbps]')

      fila_bajada_p5 = df_bajada.iloc[pos_p5 - 1]
      val_p5_baj = fila_bajada_p5['Velocidad de bajada [Kbps]']
      metricas[r]['val_p5_baj'] = val_p5_baj
      metricas[r]['res_p5_baj'] = (
          'CUMPLE' if val_p5_baj >= vel_bajada else 'NO CUMPLE'
      )

      fila_bajada_p95 = df_bajada.iloc[pos_p95 - 1]
      val_p95_baj = fila_bajada_p95['Velocidad de bajada [Kbps]']
      metricas[r]['val_p95_baj'] = val_p95_baj
      metricas[r]['res_p95_baj'] = (
          'CUMPLE' if val_p95_baj >= vel_bajada else 'NO CUMPLE'
      )

  return metricas


def ejecutar_cruce_seguro(sitios_activos, velocidad):
gray_fill = PatternFill(
      start_color='D3D3D3', end_color='D3D3D3', fill_type='solid'
  )
  white_font = Font(color='FFFFFF')

  dataframe = velocidad.copy()
  dataframe2 = sitios_activos.copy()

  # 1. SELECCIÓN DE COLUMNAS Y PREPARACIÓN
  cols_df1 = [
      'Numero de contrato',
      'Departamento',
      'Ciudad',
      'Identificador beneficiario',
      'Grupo',
      'Perfil',
      'Zona',
      'Velocidad de subida [Kbps]',
      'Velocidad de bajada [Kbps]',
      'Fecha de programacion',
      'Fecha de ejecucion',
      'Rango',
      'Duracion de la prueba [s]',
      'Dispositivo',
      'Tipo de solucion',
      'Tipo de centro digital',
      'Estado de la prueba',
      'Identificador de la prueba',
      'Centro poblado',
      'Dane institucion educativa',
      'Tipo',
  ]

  # 2. CREACIÓN DE dataframe_1 (DEBE IR AQUÍ, ANTES DEL BUCLE DE PERFILES)
  dataframe_1 = dataframe[cols_df1]
  dataframe_2 = dataframe2[['Identificador beneficiario', 'Estado']]

  dataframe_1 = pd.merge(
      dataframe_1,
      dataframe_2,
      left_on='Identificador beneficiario',
      right_on='Identificador beneficiario',
      how='inner',
  )

  # Limpieza y filtrado inicial de datos
  dataframe_1['Estado'] = dataframe_1['Estado'].astype(str).str.strip()
  dataframe_1 = dataframe_1[(dataframe_1['Estado'] == 'EN OPERACION')]
  dataframe_1 = dataframe_1[~(dataframe_1['Tipo'] == 'forzada')]
  dataframe_1['Identificador de la prueba'] = dataframe_1[
      'Identificador de la prueba'
  ].astype(str)
  dataframe_1 = dataframe_1[
      (dataframe_1['Rango'] >= 6) & (dataframe_1['Rango'] <= 20)
  ]

  # Configuración por perfil: (Nombre, vel_baj_show, vel_baj, vel_sub_show, vel_sub)
perfiles_config = {
    'PERFIL 1': ('PERFIL_1', 26.40, 26400, 6.60, 6600),
    'PERFIL 2': ('PERFIL_2', 28.50, 28500, 7.13, 7130),
    'PERFIL 3+': ('PERFIL_3+', 32.00, 32000, 8.00, 8000),
    'PERFIL 4+': ('PERFIL_4+', 41.90, 41900, 10.48, 10480),
}

for nombre_perfil, (
    tag_archivo,
    baj_show,
    vel_bajada,
    sub_show,
    vel_subida,
) in perfiles_config.items():
  # Filtrar exactamente por el string del Perfil
  df_perfil = dataframe_1[
      dataframe_1['Perfil'].astype(str).str.strip() == nombre_perfil
  ]

  # Calcular las métricas utilizando la función
  metricas = procesar_metricas_matricula(df_perfil, vel_subida, vel_bajada)

  output_path = f'/content/20260617_{tag_archivo}.xlsx'

  # Escribir primero los DataFrames en cada pestaña
  with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
    # Pestaña Principal del Perfil
    df_perfil.to_excel(
        writer, sheet_name=tag_archivo, index=False, startrow=13, startcol=1
    )

    # Pestañas de los Rangos (6 al 20)
    for r in range(6, 21):
      df_r = df_perfil[df_perfil['Rango'] == r]
      df_r.to_excel(
          writer, sheet_name=str(r), index=False, startrow=13, startcol=1
      )

    workbook = writer.book

  # Aplicar formatos y cabeceras dinámicas en openpyxl
  for sheet_name in workbook.sheetnames[1:]:
    sheet = workbook[sheet_name]
    rango = int(sheet_name)

    sheet.merge_cells('A1:V1')
    sheet['A1'] = main_title
    sheet['A1'].alignment = Alignment(horizontal='center', vertical='center')
    sheet['A1'].fill = gray_fill
    sheet['A1'].font = white_font

    for col, st1 in zip(['B'], sub_titles):
      sheet[f'{col}3'] = st1
      sheet[f'{col}3'].alignment = Alignment(
          horizontal='center', vertical='center'
      )
      sheet[f'{col}3'].fill = gray_fill
      sheet[f'{col}3'].font = white_font

    for col, st2 in zip(['B'], sub_titles2):
      sheet[f'{col}4'] = st2
      sheet[f'{col}4'].alignment = Alignment(
          horizontal='center', vertical='center'
      )

    columns = ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J']
    for col, st3 in zip(columns, sub_titles3):
      sheet[f'{col}6'] = st3
      sheet[f'{col}6'].alignment = Alignment(
          horizontal='center', vertical='center'
      )
      sheet[f'{col}6'].fill = gray_fill
      sheet[f'{col}6'].font = white_font

    # Extraer métricas reales del diccionario
    pos_p5_baj = metricas[rango]['pos_p5_baj']
    val_p5_baj = metricas[rango]['val_p5_baj']
    res_p5_baj = metricas[rango]['res_p5_baj']

    pos_p95_baj = metricas[rango]['pos_p95_baj']
    val_p95_baj = metricas[rango]['val_p95_baj']
    res_p95_baj = metricas[rango]['res_p95_baj']

    sub_titles4 = [
        'Dowload',
        tag_archivo,
        baj_show,
        pos_p5_baj,
        val_p5_baj,
        res_p5_baj,
        pos_p95_baj,
        val_p95_baj,
        res_p95_baj,
    ]
    for col, st4 in zip(columns, sub_titles4):
      sheet[f'{col}7'] = st4
      sheet[f'{col}7'].alignment = Alignment(
          horizontal='center', vertical='center'
      )

    for col, st5 in zip(columns, sub_titles5):
      sheet[f'{col}9'] = st5
      sheet[f'{col}9'].alignment = Alignment(
          horizontal='center', vertical='center'
      )
      sheet[f'{col}9'].fill = gray_fill
      sheet[f'{col}9'].font = white_font

    pos_p5_sub = metricas[rango]['pos_p5_sub']
    val_p5_sub = metricas[rango]['val_p5_sub']
    res_p5_sub = metricas[rango]['res_p5_sub']

    pos_p95_sub = metricas[rango]['pos_p95_sub']
    val_p95_sub = metricas[rango]['val_p95_sub']
    res_p95_sub = metricas[rango]['res_p95_sub']

    sub_titles6 = [
        'Upload',
        tag_archivo,
        sub_show,
        pos_p5_sub,
        val_p5_sub,
        res_p5_sub,
        pos_p95_sub,
        val_p95_sub,
        res_p95_sub,
    ]
    for col, st6 in zip(columns, sub_titles6):
      sheet[f'{col}10'] = st6
      sheet[f'{col}10'].alignment = Alignment(
          horizontal='center', vertical='center'
      )

    for row in [1, 3, 6, 9, 14]:
      sheet.row_dimensions[row].height = 60

    all_cols = [
        'A',
        'B',
        'C',
        'D',
        'E',
        'F',
        'G',
        'H',
        'I',
        'J',
        'K',
        'L',
        'M',
        'N',
        'O',
        'P',
        'Q',
        'R',
        'S',
        'T',
        'U',
        'V',
    ]
    for col in all_cols:
      sheet[f'{col}14'].alignment = Alignment(
          horizontal='center', vertical='center'
      )
      sheet[f'{col}14'].fill = gray_fill
      sheet[f'{col}14'].font = white_font

  workbook.save(output_path)
  files.download(output_path)

